# -*- coding: utf-8 -*-
"""فاز ۷ — خروجی Excel و PDF از گزارش‌ها.

- Excel: کتاب‌کار چندبرگی (راهنما، شاخص‌های کلی، گلوگاه‌ها، بار کاری، حجم روزانه)
  با شیت‌های راست‌به‌چپ، سرستون بولد، قاب، عرض ستون و قالب تاریخ جلالی
- PDF: سند فارسی راست‌به‌چپ با reportlab؛ متن فارسی با arabic_reshaper
  (اشکال نمایش) و python-bidi (ترتیب بصری) آماده می‌شود؛ قلم Vazirmatn
  همراه مخزن است و در نبود آن به FreeSans/DejaVu برمی‌گردد
"""
from __future__ import annotations

import io
import os
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font as XFont, PatternFill, Side
from openpyxl.utils import get_column_letter

import arabic_reshaper
from bidi.algorithm import get_display
from reportlab.lib import colors

import jdatetime

from app.modules.engine.models import ProcessDefinition
from app.modules.reports import service

# ---------- قلم فارسی ----------

_ASSET_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "assets", "fonts"))
_FONT_CANDIDATES = [
    (os.path.join(_ASSET_DIR, "Vazirmatn-Regular.ttf"), os.path.join(_ASSET_DIR, "Vazirmatn-Bold.ttf")),
    ("/usr/share/fonts/truetype/freefont/FreeSans.ttf", "/usr/share/fonts/truetype/freefont/FreeSansBold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]

_reshaper = arabic_reshaper.ArabicReshaper(
    configuration={"delete_harakat": False, "support_ligatures": False})
_FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def fa_digits(s) -> str:
    """تبدیل ارقام لاتین به فارسی."""
    return str(s).translate(_FA_DIGITS)


def fa(text) -> str:
    """آماده‌سازی متن فارسی برای PDF: ارقام فارسی + اشکال نمایش + ترتیب بصری."""
    s = fa_digits(text)
    s = get_display(_reshaper.reshape(s), base_dir="R")
    return s


def jalali(iso_date: str) -> str:
    """تبدیل تاریخ میلادی ISO (YYYY-MM-DD) به جلالی YYYY/MM/DD."""
    try:
        d = datetime.fromisoformat(str(iso_date)[:10]).date()
        return jdatetime.date.fromgregorian(date=d).strftime("%Y/%m/%d")
    except (ValueError, TypeError):
        return str(iso_date)


def fmt_dur(sec: float | None) -> str:
    """ثانیه → عبارت خوانا با ارقام فارسی."""
    if sec is None or sec != sec:  # NaN
        return "—"
    if sec <= 0:
        return "۰"
    if sec < 60:
        return f"{fa_digits(round(sec))} ثانیه"
    if sec < 3600:
        return f"{fa_digits(round(sec / 60, 1))} دقیقه"
    if sec < 86400:
        return f"{fa_digits(round(sec / 3600, 1))} ساعت"
    return f"{fa_digits(round(sec / 86400, 1))} روز"


STATUS_FA = {"running": "در جریان", "completed": "تکمیل‌شده", "cancelled": "لغوشده", "error": "خطا"}
DATASET_FA = {"overview": "شاخص‌های کلی", "bottlenecks": "گلوگاه‌ها",
              "workload": "بار کاری", "volume": "حجم روزانه"}


# ---------- گردآوری داده ----------

def collect_report(db, days: int = 30, date_from: str | None = None,
                   date_to: str | None = None, definition_id: int | None = None,
                   status: str | None = None, dataset: str | None = None) -> dict:
    """گردآوری داده گزارش (همه بخش‌ها یا فقط یک مجموعه‌داده) + فراداده."""
    common = dict(days=days, date_from=date_from, date_to=date_to, definition_id=definition_id)
    data: dict = {
        "meta": {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "days": days, "date_from": date_from or "", "date_to": date_to or "",
            "definition_id": definition_id, "status": status or "",
            "definition_title": "", "dataset": dataset or "all",
        },
        "overview": None, "bottlenecks": None, "workload": None, "volume": None,
    }
    if definition_id:
        pd = db.get(ProcessDefinition, definition_id)
        if pd:
            data["meta"]["definition_title"] = pd.title
    want = (lambda k: dataset in (None, "all") or k == dataset)
    if want("overview"):
        data["overview"] = service.kpi_overview(db, status=status, **common)
    if want("bottlenecks"):
        data["bottlenecks"] = service.kpi_bottlenecks(db, **common)
    if want("workload"):
        data["workload"] = service.kpi_workload(db, **common)
    if want("volume"):
        data["volume"] = service.kpi_volume(db, **common)
    return data


# ---------- Excel ----------

_TH_FILL = PatternFill("solid", fgColor="1E3A8A")
_TH_FONT = XFont(name="Vazirmatn", bold=True, color="FFFFFF", size=11)
_BODY_FONT = XFont(name="Vazirmatn", size=10)
_TITLE_FONT = XFont(name="Vazirmatn", bold=True, size=13, color="1E3A8A")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_R = Alignment(horizontal="right", vertical="center", wrap_text=True)
_C = Alignment(horizontal="center", vertical="center")


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]) -> tuple:
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font, cell.border, cell.alignment = _TH_FILL, _TH_FONT, _BORDER, _C
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    return ws, len(headers)


def _row(ws, values: list, numeric_cols: set[int] | None = None):
    numeric_cols = numeric_cols or set()
    ws.append(values)
    r = ws.max_row
    for i, _ in enumerate(values, start=1):
        cell = ws.cell(row=r, column=i)
        cell.font, cell.border = _BODY_FONT, _BORDER
        cell.alignment = _C if i in numeric_cols else _R


def _kv(ws, label: str, value):
    ws.append([label, value])
    r = ws.max_row
    ws.cell(row=r, column=1).font = XFont(name="Vazirmatn", bold=True, size=10, color="334155")
    ws.cell(row=r, column=1).alignment = _R
    c2 = ws.cell(row=r, column=2)
    c2.font, c2.alignment = _BODY_FONT, _C
    ws.cell(row=r, column=1).border = _BORDER
    c2.border = _BORDER


def excel_bytes(data: dict, report_name: str = "گزارش KPI سامانه BPMS") -> bytes:
    """ساخت فایل xlsx از داده گزارش."""
    meta = data["meta"]
    wb = Workbook()
    wb.remove(wb.active)

    # --- شیت راهنما / خلاصه ---
    ws = wb.create_sheet("راهنما")
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 46
    ws.append([report_name])
    ws.cell(row=1, column=1).font = _TITLE_FONT
    ws.append([])
    info = [
        ("زمان تولید", fa_digits(datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")) + " (UTC)"),
        ("بازه گزارش", f"{fa_digits(meta.get('days') or 30)} روز اخیر"
         + (f" | از {jalali(meta['date_from'][:10])} تا {jalali(meta['date_to'][:10])}"
            if meta.get("date_from") else "")),
        ("فرایند", meta.get("definition_title") or "همه فرایندها"),
        ("وضعیت نمونه‌ها", STATUS_FA.get(meta.get("status"), "همه وضعیت‌ها")),
        ("مجموعه‌داده‌ها", "، ".join(DATASET_FA[k] for k in ("overview", "bottlenecks", "workload", "volume")
                                    if data.get(k) is not None)),
        ("سامانه", "BPMS — سامانه مدیریت فرایندهای کسب‌وکار (فاز ۷)"),
    ]
    for label, value in info:
        _kv(ws, label, value)

    # --- شاخص‌های کلی ---
    if data.get("overview") is not None:
        ov = data["overview"]
        ws, _ = _sheet(wb, "شاخص‌های کلی", ["شاخص", "مقدار"], [36, 30])
        ins, cyc, tsk, sla = ov["instances"], ov["cycle_time_seconds"], ov["tasks"], ov["sla"]
        rows = [
            ("مجموع نمونه‌ها", fa_digits(ins["total"])),
            ("در جریان", fa_digits(ins["by_status"].get("running", 0))),
            ("تکمیل‌شده", fa_digits(ins["by_status"].get("completed", 0))),
            ("لغوشده", fa_digits(ins["by_status"].get("cancelled", 0))),
            ("خطا", fa_digits(ins["by_status"].get("error", 0))),
            ("نرخ تکمیل", f"{fa_digits(ins['completion_rate'])}٪"),
            ("میانگین نمونه در روز", fa_digits(ins["avg_per_day"])),
            ("میانگین زمان چرخه", fmt_dur(cyc["avg"])),
            ("کمینه زمان چرخه", fmt_dur(cyc["min"])),
            ("بیشینه زمان چرخه", fmt_dur(cyc["max"])),
            ("صدک ۵۰ چرخه", fmt_dur(cyc["p50"])),
            ("صدک ۹۵ چرخه", fmt_dur(cyc["p95"])),
            ("مجموع وظایف", fa_digits(tsk["total"])),
            ("وظایف آماده", fa_digits(tsk["ready"])),
            ("وظایف در دست اقدام", fa_digits(tsk["claimed"])),
            ("وظایف تکمیل‌شده", fa_digits(tsk["completed"])),
            ("سرریز فعال", fa_digits(tsk["overdue_active"])),
            ("میانگین انتظار وظیفه", fmt_dur(tsk["avg_wait_seconds"])),
            ("میانگین رسیدگی وظیفه", fmt_dur(tsk["avg_handling_seconds"])),
            ("هشدار SLA", fa_digits(sla["warned"])),
            ("فرازدگی SLA", fa_digits(sla["escalated"])),
            ("نرخ فرازدگی SLA", f"{fa_digits(sla['breach_rate'])}٪"),
        ]
        for label, value in rows:
            _kv(ws, label, value)

    # --- گلوگاه‌ها ---
    if data.get("bottlenecks") is not None:
        ws, _ = _sheet(wb, "گلوگاه‌ها",
                       ["گام فرایند", "شناسه گام", "ایجاد", "تکمیل", "در انتظار",
                        "میانگین انتظار", "میانگین رسیدگی", "بیشترین", "صدک ۹۵", "سرریز فعال", "فرازدگی SLA"],
                       [22, 18, 9, 9, 10, 14, 14, 13, 13, 11, 12])
        for r_ in data["bottlenecks"]:
            _row(ws, [
                r_["name"], r_["spec_id"], r_["created"], r_["completed"], r_["pending"],
                fmt_dur(r_["avg_wait_seconds"]), fmt_dur(r_["avg_handling_seconds"]),
                fmt_dur(r_["max_handling_seconds"]), fmt_dur(r_["p95_handling_seconds"]),
                r_["overdue_active"], r_["sla_escalated"],
            ], numeric_cols=set(range(3, 12)))

    # --- بار کاری ---
    if data.get("workload") is not None:
        ws, _ = _sheet(wb, "بار کاری",
                       ["کاربر / صف", "نوع", "تکمیل‌شده", "فعال", "سرریز فعال",
                        "فرازدگی SLA", "میانگین رسیدگی"],
                       [28, 12, 12, 10, 12, 13, 15])
        for r_ in data["workload"]:
            _row(ws, [
                r_["label"], "نقش" if r_["kind"] == "role" else "کاربر",
                r_["completed"], r_["active"], r_["overdue_active"], r_["sla_escalated"],
                fmt_dur(r_["avg_handling_seconds"]),
            ], numeric_cols=set(range(3, 8)))

    # --- حجم روزانه ---
    if data.get("volume") is not None:
        ws, _ = _sheet(wb, "حجم روزانه", ["تاریخ (جلالی)", "تاریخ (میلادی)", "شروع‌شده", "تکمیل‌شده"],
                       [16, 16, 11, 12])
        for r_ in data["volume"]:
            if not (r_["started"] or r_["completed"]):
                continue
            _row(ws, [jalali(r_["date"]), r_["date"], r_["started"], r_["completed"]],
                 numeric_cols={3, 4})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

_PDF_STYLES: dict | None = None
_PDF_FONTS_OK = False


def _register_pdf_fonts() -> None:
    """ثبت قلم فارسی (یک‌بار)."""
    global _PDF_STYLES, _PDF_FONTS_OK
    if _PDF_STYLES is not None:
        return
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for reg, bold in _FONT_CANDIDATES:
        if os.path.exists(reg) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("BPMS-Fa", reg))
            pdfmetrics.registerFont(TTFont("BPMS-Fa-Bold", bold))
            break
    else:  # pragma: no cover — قلم پیش‌فرض Helvetica (فارسی رندر نمی‌شود)
        pdfmetrics.registerFont(TTFont("BPMS-Fa", "Helvetica"))
        pdfmetrics.registerFont(TTFont("BPMS-Fa-Bold", "Helvetica-Bold"))

    base = dict(fontName="BPMS-Fa", fontSize=9, leading=14, alignment=TA_RIGHT)
    _PDF_STYLES = {
        "title": ParagraphStyle("t", fontName="BPMS-Fa-Bold", fontSize=16, leading=22,
                                alignment=TA_CENTER, textColor=colors.HexColor("#1E3A8A")),
        "meta": ParagraphStyle("m", **{**base, "fontSize": 8.5,
                                       "textColor": colors.HexColor("#475569"),
                                       "alignment": TA_CENTER}),
        "h2": ParagraphStyle("h2", fontName="BPMS-Fa-Bold", fontSize=11.5, leading=17,
                             alignment=TA_RIGHT, textColor=colors.HexColor("#1E3A8A"),
                             spaceBefore=10, spaceAfter=4),
        "cell": ParagraphStyle("c", **base),
        "cellb": ParagraphStyle("cb", fontName="BPMS-Fa-Bold", fontSize=9, leading=14,
                                alignment=TA_RIGHT, textColor=colors.HexColor("#0F172A")),
        "empty": ParagraphStyle("e", **{**base, "fontSize": 9,
                                        "textColor": colors.HexColor("#64748B"),
                                        "alignment": TA_CENTER}),
    }
    _PDF_FONTS_OK = True


def _P(text: str, style: str = "cell"):
    from reportlab.platypus import Paragraph
    return Paragraph(fa(text), _PDF_STYLES[style])


def _table(headers: list[str], rows: list[list], widths: list[float], bold_first: bool = False):
    """جدول RTL — ترتیب ستون‌ها معکوس می‌شود تا ستون نخست سمت راست باشد."""
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    body = [[_P(h, "cellb") for h in reversed(headers)]]
    for r in rows:
        cells = []
        for j, v in enumerate(reversed(r)):
            st = "cellb" if (bold_first and j == len(r) - 1) else "cell"
            cells.append(_P(str(v), st))
        body.append(cells)
    t = Table(body, colWidths=list(reversed(widths)), repeatRows=1, hAlign="CENTER")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _kv_table(pairs: list[tuple[str, str]]):
    """جدول دوستونی شاخص/مقدار در چهار ستون (دو جفت در هر ردیف)."""
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    body = []
    for i in range(0, len(pairs), 2):
        chunk = pairs[i:i + 2]
        row = []
        for label, value in reversed(chunk):  # جفت نخست سمت راست
            row.extend([_P(label, "cellb"), _P(value, "cell")])
        if len(chunk) == 1:
            row.extend(["", ""])
        body.append(row)
    t = Table(body, colWidths=[95, 90, 95, 90], hAlign="CENTER")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1F5F9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("BPMS-Fa", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawCentredString(doc.pagesize[0] / 2, 8 * 2.83 * 2, fa(f"صفحه {fa_digits(canvas.getPageNumber())}"))
    canvas.drawRightString(doc.pagesize[0] - 15 * 2.83, 8 * 2.83 * 2, fa("سامانه BPMS — گزارش‌ساز و KPI"))
    canvas.restoreState()


def pdf_bytes(data: dict, report_name: str = "گزارش KPI سامانه BPMS") -> bytes:
    """ساخت PDF فارسی راست‌به‌چپ از داده گزارش."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Spacer

    _register_pdf_fonts()
    meta = data["meta"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=14 * mm, bottomMargin=16 * mm,
        leftMargin=14 * mm, rightMargin=14 * mm, title=report_name,
        author="BPMS")
    story = []

    story.append(_P(report_name, "title"))
    period = f"{fa_digits(meta.get('days') or 30)} روز اخیر"
    if meta.get("date_from"):
        period = f"{jalali(meta['date_from'][:10])} تا {jalali(meta['date_to'][:10])}"
    scope = "، ".join(x for x in [
        meta.get("definition_title") or "همه فرایندها",
        f"وضعیت: {STATUS_FA.get(meta.get('status'), 'همه')}" if meta.get("status") else "",
    ] if x)
    gen = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    story.append(Spacer(1, 3))
    story.append(_P(f"بازه: {period}" + (f" | {scope}" if scope else ""), "meta"))
    story.append(_P(f"زمان تولید: {fa_digits(gen)} (UTC)", "meta"))
    story.append(Spacer(1, 6))

    # --- شاخص‌های کلی ---
    if data.get("overview"):
        ov = data["overview"]
        ins, cyc, tsk, sla = ov["instances"], ov["cycle_time_seconds"], ov["tasks"], ov["sla"]
        story.append(_P("شاخص‌های کلی", "h2"))
        story.append(_kv_table([
            ("مجموع نمونه‌ها", fa_digits(ins["total"])),
            ("نرخ تکمیل", f"{fa_digits(ins['completion_rate'])}٪"),
            ("در جریان", fa_digits(ins["by_status"].get("running", 0))),
            ("تکمیل‌شده", fa_digits(ins["by_status"].get("completed", 0))),
            ("میانگین زمان چرخه", fmt_dur(cyc["avg"])),
            ("صدک ۹۵ چرخه", fmt_dur(cyc["p95"])),
            ("مجموع وظایف", fa_digits(tsk["total"])),
            ("سرریز فعال", fa_digits(tsk["overdue_active"])),
            ("میانگین انتظار وظیفه", fmt_dur(tsk["avg_wait_seconds"])),
            ("میانگین رسیدگی وظیفه", fmt_dur(tsk["avg_handling_seconds"])),
            ("هشدار SLA", fa_digits(sla["warned"])),
            ("فرازدگی SLA", f"{fa_digits(sla['escalated'])} ({fa_digits(sla['breach_rate'])}٪)"),
        ]))
        story.append(Spacer(1, 4))

    # --- گلوگاه‌ها ---
    if data.get("bottlenecks") is not None:
        story.append(_P("گلوگاه‌های فرایند (به تفکیک گام)", "h2"))
        rows = data["bottlenecks"]
        if rows:
            story.append(_table(
                ["گام فرایند", "ایجاد", "تکمیل", "در انتظار", "میانگین انتظار",
                 "میانگین رسیدگی", "بیشترین", "صدک ۹۵", "سرریز SLA"],
                [[r["name"], fa_digits(r["created"]), fa_digits(r["completed"]),
                  fa_digits(r["pending"]), fmt_dur(r["avg_wait_seconds"]),
                  fmt_dur(r["avg_handling_seconds"]), fmt_dur(r["max_handling_seconds"]),
                  fmt_dur(r["p95_handling_seconds"]), fa_digits(r["sla_escalated"])] for r in rows],
                [88, 32, 34, 36, 55, 58, 52, 52, 42], bold_first=True))
        else:
            story.append(_P("داده‌ای در این بازه نیست", "empty"))

    # --- بار کاری ---
    if data.get("workload") is not None:
        story.append(_P("بار کاری افراد و صف‌ها", "h2"))
        rows = data["workload"]
        if rows:
            story.append(_table(
                ["کاربر / صف", "نوع", "تکمیل‌شده", "فعال", "سرریز فعال", "فرازدگی SLA", "میانگین رسیدگی"],
                [[r["label"], "نقش" if r["kind"] == "role" else "کاربر", fa_digits(r["completed"]),
                  fa_digits(r["active"]), fa_digits(r["overdue_active"]),
                  fa_digits(r["sla_escalated"]), fmt_dur(r["avg_handling_seconds"])] for r in rows],
                [95, 34, 40, 34, 42, 46, 60], bold_first=True))
        else:
            story.append(_P("داده‌ای در این بازه نیست", "empty"))

    # --- حجم روزانه (روزهای دارای فعالیت) ---
    if data.get("volume") is not None:
        story.append(_P("حجم روزانه نمونه‌ها (فقط روزهای دارای فعالیت)", "h2"))
        rows = [r for r in data["volume"] if (r["started"] or r["completed"])]
        if rows:
            story.append(_table(
                ["تاریخ (جلالی)", "شروع‌شده", "تکمیل‌شده"],
                [[jalali(r["date"]), fa_digits(r["started"]), fa_digits(r["completed"])] for r in rows],
                [140, 90, 90]))
        else:
            story.append(_P("داده‌ای در این بازه نیست", "empty"))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()

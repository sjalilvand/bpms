# -*- coding: utf-8 -*-
"""آزمون‌های فاز ۷ — خروجی Excel و PDF از گزارش‌ها.

- اعتبارسنجی بایت‌های خروجی (امضای PK برای xlsx و %PDF برای pdf)
- ساختار کتاب‌کار Excel (برگه‌ها، سرستون، راست‌به‌چپ)
- سند PDF (حجم و ساختار) و خروجی گزارش ذخیره‌شده
- مجوزها: report.view لازم است؛ کاربر بدون مجوز → ۴۰۳
"""
import io
import time
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.database import SessionLocal
from app.main import app
from app.modules.engine.models import EngineTask, ProcessDefinition, ProcessInstance, ProcessVersion
from app.modules.identity.models import User
from app.modules.reports.models import SavedReport

BKEY = "export-test-marker"


@pytest.fixture(scope="module")
def api():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def admin_token(api):
    r = api.post("/api/v1/auth/login", json={"username": "admin", "password": "Bpms@1234"})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def _h(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def seeded():
    """داده آزمون خروجی — یک فرایند با دو نمونه و سه وظیفه."""
    db = SessionLocal()
    try:
        admin = db.scalar(select(User).where(User.username == "admin"))
        now = datetime.now(timezone.utc).replace(tzinfo=None)

        pd = ProcessDefinition(key=f"exp_proc_{int(time.time())}", title="فرایند آزمون خروجی")
        db.add(pd)
        db.flush()
        pv = ProcessVersion(definition_id=pd.id, version=1, bpmn_xml="<x/>",
                            main_process_id="exp", config_json={})
        db.add(pv)
        db.flush()

        i1 = ProcessInstance(definition_id=pd.id, version_id=pv.id, business_key=BKEY,
                             status="completed", started_by_id=admin.id,
                             started_at=now - timedelta(days=2),
                             finished_at=now - timedelta(days=1))
        db.add(i1)
        db.flush()
        t1 = EngineTask(instance_id=i1.id, task_uuid="x-e1", spec_id="Task_Review",
                        name="بررسی", status="completed", assignee_id=admin.id,
                        created_at=now - timedelta(days=2),
                        claimed_at=now - timedelta(days=2) + timedelta(hours=1),
                        completed_at=now - timedelta(days=2) + timedelta(hours=3))
        i2 = ProcessInstance(definition_id=pd.id, version_id=pv.id, business_key=BKEY,
                             status="running", started_by_id=admin.id,
                             started_at=now - timedelta(hours=5))
        db.add(i2)
        db.flush()
        t2 = EngineTask(instance_id=i2.id, task_uuid="x-e2", spec_id="Task_Review",
                        name="بررسی", status="ready", created_at=now - timedelta(hours=5),
                        due_at=now - timedelta(hours=1))
        t3 = EngineTask(instance_id=i2.id, task_uuid="x-e3", spec_id="Task_Approve",
                        name="تأیید", status="ready", created_at=now - timedelta(hours=4))
        db.add_all([t1, t2, t3])
        db.commit()
        out = {"definition_id": pd.id, "version_id": pv.id,
               "instance_ids": [i1.id, i2.id], "task_ids": [t1.id, t2.id, t3.id]}
        yield out
    finally:
        # پاکسازی پس از ماژول
        try:
            ids = out["task_ids"]
            for t in db.scalars(select(EngineTask).where(EngineTask.id.in_(ids))).all():
                db.delete(t)
            for it in db.scalars(select(ProcessInstance).where(
                    ProcessInstance.id.in_(out["instance_ids"]))).all():
                db.delete(it)
            for v in db.scalars(select(ProcessVersion).where(
                    ProcessVersion.id == out["version_id"])).all():
                db.delete(v)
            for p in db.scalars(select(ProcessDefinition).where(
                    ProcessDefinition.id == out["definition_id"])).all():
                db.delete(p)
            for sr in db.scalars(select(SavedReport).where(
                    SavedReport.name.like("خروجی آزمون%"))).all():
                db.delete(sr)
            db.commit()
        except Exception:
            db.rollback()
        finally:
            db.close()


class TestExportExcel:
    def test_excel_export_all_datasets(self, api, admin_token, seeded):
        r = api.get("/api/v1/reports/export/excel",
                    params={"days": 30, "definition_id": seeded["definition_id"]},
                    headers=_h(admin_token))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert "attachment" in r.headers.get("content-disposition", "")
        body = r.content
        assert body[:2] == b"PK"          # امضای zip/OOXML
        assert len(body) > 2000
        # ساختار کتاب‌کار
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(body))
        names = wb.sheetnames
        for expected in ("راهنما", "شاخص‌های کلی", "گلوگاه‌ها", "بار کاری", "حجم روزانه"):
            assert expected in names, names
        # سرستون برگه گلوگاه‌ها
        ws = wb["گلوگاه‌ها"]
        assert ws.cell(row=1, column=1).value == "گام فرایند"
        assert ws.sheet_view.rightToLeft is True
        # ردیف داده گلوگاه
        assert ws.cell(row=2, column=1).value == "بررسی"

    def test_excel_export_empty_range(self, api, admin_token):
        r = api.get("/api/v1/reports/export/excel",
                    params={"date_from": "2001-01-01", "date_to": "2001-01-02"},
                    headers=_h(admin_token))
        assert r.status_code == 200
        assert r.content[:2] == b"PK"


class TestExportPdf:
    def test_pdf_export(self, api, admin_token, seeded):
        r = api.get("/api/v1/reports/export/pdf",
                    params={"days": 30, "definition_id": seeded["definition_id"]},
                    headers=_h(admin_token))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/pdf")
        body = r.content
        assert body[:5] == b"%PDF-"
        assert len(body) > 5000  # قلم جاسازی‌شده حجم قابل توجهی دارد

    def test_pdf_export_empty_range(self, api, admin_token):
        r = api.get("/api/v1/reports/export/pdf",
                    params={"date_from": "2001-01-01", "date_to": "2001-01-02"},
                    headers=_h(admin_token))
        assert r.status_code == 200
        assert r.content[:5] == b"%PDF-"


class TestSavedExport:
    def test_saved_report_export_excel_and_pdf(self, api, admin_token, seeded):
        h = _h(admin_token)
        r = api.post("/api/v1/reports/saved", headers=h, json={
            "name": "خروجی آزمون — گلوگاه‌ها", "description": "آزمون خروجی",
            "dataset": "bottlenecks",
            "filters": {"days": 30, "definition_id": seeded["definition_id"]}})
        assert r.status_code == 200, r.text
        rid = r.json()["id"]

        rx = api.get(f"/api/v1/reports/saved/{rid}/export",
                     params={"format": "excel"}, headers=h)
        assert rx.status_code == 200, rx.text
        assert rx.content[:2] == b"PK"
        assert rx.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        rp = api.get(f"/api/v1/reports/saved/{rid}/export",
                     params={"format": "pdf"}, headers=h)
        assert rp.status_code == 200, rp.text
        assert rp.content[:5] == b"%PDF-"

    def test_saved_export_unknown_report(self, api, admin_token):
        r = api.get("/api/v1/reports/saved/999999/export",
                    params={"format": "excel"}, headers=_h(admin_token))
        assert r.status_code == 404

    def test_saved_export_invalid_format(self, api, admin_token):
        r = api.get("/api/v1/reports/saved/1/export",
                    params={"format": "docx"}, headers=_h(admin_token))
        assert r.status_code == 422


class TestExportPermissions:
    def test_export_requires_report_view(self, api):
        # کاربر operator بدون report.view → ۴۰۳
        r = api.post("/api/v1/auth/login", json={"username": "user1", "password": "Bpms@1234"})
        token = r.json()["access_token"]
        r1 = api.get("/api/v1/reports/export/excel", headers=_h(token))
        r2 = api.get("/api/v1/reports/export/pdf", headers=_h(token))
        assert r1.status_code == 403
        assert r2.status_code == 403

    def test_export_requires_auth(self, api):
        assert api.get("/api/v1/reports/export/excel").status_code == 401
        assert api.get("/api/v1/reports/export/pdf").status_code == 401


class TestExportHelpers:
    def test_fa_digits_and_duration(self):
        from app.modules.reports.exporting import fa_digits, fmt_dur, jalali
        assert fa_digits("123") == "۱۲۳"
        assert fmt_dur(7200) == "۲.۰ ساعت"
        assert fmt_dur(45) == "۴۵ ثانیه"
        assert fmt_dur(0) == "۰"
        assert fmt_dur(None) == "—"
        assert len(jalali("2026-10-01").split("/")) == 3

    def test_fa_visual_order(self):
        from app.modules.reports.exporting import fa
        out = fa("گزارش KPI")
        assert out  # خروجی بصری غیرخالی
        assert "KPI" in out  # لاتین حفظ می‌شود
